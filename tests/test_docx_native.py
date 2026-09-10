"""Tests for bibr.input.docx_native.DocxParser."""

import io

import pytest

pytest.importorskip("docx")

from docx import Document

from bibr.input.docx_native import _NS, DocxParser
from bibr.paper_contents import CanonicalSection

_NS_W = _NS["w"]


def _make_docx_bytes(build_fn) -> bytes:
    """Helper: build a python-docx Document via callback, return DOCX bytes."""
    doc = Document()
    build_fn(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestDocxParserHeadings:
    def test_title_and_headings_become_sections(self):
        def build(doc):
            doc.add_heading("My Paper Title", level=0)  # Title style
            doc.add_heading("Introduction", level=1)
            doc.add_paragraph("Intro body text.")
            doc.add_heading("Methods", level=1)
            doc.add_paragraph("Method body text.")

        contents = _parse(build)
        headers = [s.header for s in contents.sections if s.level > 0]
        assert "My Paper Title" in headers
        assert "Introduction" in headers
        assert "Methods" in headers

    def test_heading_levels_preserved(self):
        def build(doc):
            doc.add_heading("Methods", level=1)
            doc.add_heading("Participants", level=2)
            doc.add_heading("Eligibility", level=3)

        contents = _parse(build)
        levels = {s.header: s.level for s in contents.sections if s.level > 0}
        assert levels["Methods"] == 1
        assert levels["Participants"] == 2
        assert levels["Eligibility"] == 3

    def test_detected_title_from_heading_one_when_no_title_style(self):
        def build(doc):
            # No Title style — first Heading 1 should be picked up
            doc.add_heading("Paper About X", level=1)
            doc.add_paragraph("Body text.")

        contents = _parse(build)
        assert contents.detected_title == "Paper About X"

    def test_numbering_deepens_level(self):
        def build(doc):
            # All paragraph_title-equivalent (Heading 2 in DOCX); numbering pushes deeper
            doc.add_heading("Methods", level=2)
            doc.add_heading("2.3.1 Sub-procedure", level=2)

        contents = _parse(build)
        sub = next(s for s in contents.sections if s.header.startswith("2.3.1"))
        assert sub.level == 3


class TestDocxParserBody:
    def test_body_paragraphs_become_sentences(self):
        def build(doc):
            doc.add_heading("Introduction", level=1)
            doc.add_paragraph("This is the first sentence. This is the second sentence.")

        parser = DocxParser(_make_docx_bytes(build))
        contents = parser.parse()
        # Pre-segmentation: deferred entries exist, sentences empty
        assert len(parser._deferred_texts) == 1
        assert contents.sentences == []
        # Apply trivial single-segment fallback
        parser.apply_segmentation(
            contents,
            [["This is the first sentence.", "This is the second sentence."]],
        )
        assert len(contents.sentences) == 2
        assert contents.sentences[0].text.startswith("This is the first")

    def test_section_assignment(self):
        def build(doc):
            doc.add_heading("Methods", level=1)
            doc.add_paragraph("Methods body.")

        parser = DocxParser(_make_docx_bytes(build))
        contents = parser.parse()
        parser.apply_segmentation(contents, [["Methods body."]])
        methods_sec = next(s for s in contents.sections if s.header == "Methods")
        assert contents.sentences[0].section_id == methods_sec.section_id


class TestDocxParserTables:
    def test_table_extracted(self):
        def build(doc):
            doc.add_heading("Results", level=1)
            t = doc.add_table(rows=2, cols=2)
            t.cell(0, 0).text = "A"
            t.cell(0, 1).text = "B"
            t.cell(1, 0).text = "1"
            t.cell(1, 1).text = "2"

        contents = _parse(build)
        assert len(contents.tables) == 1
        tbl = contents.tables[0]
        assert tbl.df.columns.tolist() == ["A", "B"]
        assert tbl.df.iloc[0].tolist() == ["1", "2"]


class TestDocxParserUrls:
    def test_url_in_body_detected(self):
        def build(doc):
            doc.add_heading("Methods", level=1)
            doc.add_paragraph("See https://example.com/data for more.")

        parser = DocxParser(_make_docx_bytes(build))
        contents = parser.parse()
        parser.apply_segmentation(contents, [["See https://example.com/data for more."]])
        urls = [link.url for link in contents.links]
        assert "https://example.com/data" in urls


class TestDocxParserContentSections:
    def test_create_content_sections_adds_table_section(self):
        def build(doc):
            doc.add_heading("Results", level=1)
            t = doc.add_table(rows=1, cols=1)
            t.cell(0, 0).text = "X"

        parser = DocxParser(_make_docx_bytes(build))
        contents = parser.parse()
        parser.apply_segmentation(contents, [])  # no body paragraphs
        parser.create_content_sections(contents)
        table_sections = [s for s in contents.sections if s.section_type == CanonicalSection.TABLE]
        assert len(table_sections) == 1
        assert table_sections[0].header == "Table 1"


def _parse(build_fn):
    """Convenience: build DOCX, parse, apply trivial segmentation per deferred entry."""
    parser = DocxParser(_make_docx_bytes(build_fn))
    contents = parser.parse()
    n_seg = sum(1 for _, _, _, needs_seg, _ in parser._deferred_texts if needs_seg)
    parser.apply_segmentation(contents, [[] for _ in range(n_seg)])
    return contents


# ----- Footnote helper -----


class TestLoadFootnotes:
    def test_returns_empty_when_no_footnotes_part(self):
        from bibr.input.docx_native import _load_footnotes

        doc_bytes = _make_docx_bytes(lambda d: d.add_paragraph("Plain body."))
        doc = Document(io.BytesIO(doc_bytes))
        # python-docx may or may not expose footnotes_part on a no-footnotes doc
        assert _load_footnotes(doc) == {}

    def test_parses_footnote_text_and_skips_separators(self):
        """Synthesize a footnotes_part on a stub doc."""
        from lxml import etree

        from bibr.input.docx_native import _NS, _load_footnotes

        w = _NS["w"]
        xml = f"""
        <w:footnotes xmlns:w="{w}">
          <w:footnote w:type="separator" w:id="-1"/>
          <w:footnote w:type="continuationSeparator" w:id="0"/>
          <w:footnote w:id="1">
            <w:p><w:r><w:t>First footnote text.</w:t></w:r></w:p>
          </w:footnote>
          <w:footnote w:id="2">
            <w:p><w:r><w:t>Second </w:t></w:r><w:r><w:t>footnote.</w:t></w:r></w:p>
          </w:footnote>
        </w:footnotes>
        """
        root = etree.fromstring(xml.strip())

        class _FakePart:
            element = root

        class _FakeDoc:
            class part:
                footnotes_part = _FakePart()

        out = _load_footnotes(_FakeDoc())
        assert out == {"1": "First footnote text.", "2": "Second footnote."}


# ----- OMML helper -----


class TestOmmlToText:
    def test_extracts_concatenated_math_text(self):
        from lxml import etree

        from bibr.input.docx_native import _NS, _omml_to_text

        m = _NS["m"]
        xml = f"""
        <m:oMath xmlns:m="{m}">
          <m:r><m:t>x</m:t></m:r>
          <m:r><m:t>+</m:t></m:r>
          <m:r><m:t>y</m:t></m:r>
        </m:oMath>
        """
        el = etree.fromstring(xml.strip())
        assert _omml_to_text(el) == "x+y"


# ----- Inline images -----


def _png_bytes() -> bytes:
    """Tiny 1x1 PNG (red pixel)."""
    import base64

    # 1x1 red PNG
    return base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )


class TestDocxParserImages:
    def test_inline_image_creates_figure(self):
        def build(doc):
            doc.add_heading("Results", level=1)
            doc.add_picture(io.BytesIO(_png_bytes()))
            doc.add_paragraph("Figure 1: A red dot.", style="Caption")

        parser = DocxParser(_make_docx_bytes(build))
        contents = parser.parse()
        # No body sentences after caption is consumed, so no segments needed
        n_seg = sum(1 for _, _, _, ns, _ in parser._deferred_texts if ns)
        parser.apply_segmentation(contents, [[] for _ in range(n_seg)])
        parser.create_content_sections(contents)

        assert len(contents.figures) == 1
        fig = contents.figures[0]
        assert fig.caption == "Figure 1: A red dot."
        assert fig.image_b64 is not None and len(fig.image_b64) > 0

        from bibr.paper_contents import CanonicalSection

        fig_sections = [s for s in contents.sections if s.section_type == CanonicalSection.FIGURE]
        assert len(fig_sections) == 1
        assert fig_sections[0].header == "Figure 1"


# ----- Display math -----


class TestDocxParserDisplayMath:
    def test_omath_para_emitted_as_formula_deferred(self):
        from docx.oxml.ns import qn
        from lxml import etree

        from bibr.input.docx_native import _NS

        m = _NS["m"]
        # Build a doc, then inject an oMathPara block after a heading
        doc = Document()
        doc.add_heading("Methods", level=1)
        body = doc.element.body
        # Construct <m:oMathPara><m:oMath><m:r><m:t>E=mc^2</m:t></m:r></m:oMath></m:oMathPara>
        omath_para_xml = f"""
        <m:oMathPara xmlns:m="{m}">
          <m:oMath><m:r><m:t>E=mc^2</m:t></m:r></m:oMath>
        </m:oMathPara>
        """
        omath_para = etree.fromstring(omath_para_xml.strip())
        # Insert before sectPr (last child)
        sectpr = body.find(qn("w:sectPr"))
        if sectpr is not None:
            sectpr.addprevious(omath_para)
        else:
            body.append(omath_para)

        buf = io.BytesIO()
        doc.save(buf)
        parser = DocxParser(buf.getvalue())
        parser.parse()
        formula_entries = [t for t in parser._deferred_texts if t[4]]  # is_formula
        assert any("E=mc^2" in entry[0] for entry in formula_entries)
        assert all(entry[0].startswith("$$") for entry in formula_entries)


# ----- Footnotes integration -----


class TestDocxParserFootnotes:
    def test_footnote_creates_section_and_xref(self):
        """Inject a footnoteReference + footnotes_part directly via XML."""
        from docx.opc.constants import CONTENT_TYPE as CT
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.opc.packuri import PackURI
        from docx.oxml.ns import qn
        from docx.parts.story import StoryPart
        from lxml import etree

        from bibr.input.docx_native import _NS

        w = _NS["w"]

        doc = Document()
        doc.add_heading("Intro", level=1)
        para = doc.add_paragraph("Body text with footnote.")
        # Insert a footnoteReference inside a new run at the end of para
        run = para.add_run()
        ref = etree.SubElement(run._r, qn("w:footnoteReference"))
        ref.set(qn("w:id"), "1")

        # Add a footnotes part with one real footnote (id=1)
        footnotes_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
        <w:footnotes xmlns:w="{w}">
          <w:footnote w:id="1">
            <w:p><w:r><w:t>This is a footnote.</w:t></w:r></w:p>
          </w:footnote>
        </w:footnotes>"""
        partname = PackURI("/word/footnotes.xml")
        fn_part = StoryPart(
            partname,
            CT.WML_FOOTNOTES,
            etree.fromstring(footnotes_xml.strip().encode("utf-8")),
            doc.part.package,
        )
        doc.part.relate_to(fn_part, RT.FOOTNOTES)

        buf = io.BytesIO()
        doc.save(buf)

        parser = DocxParser(buf.getvalue())
        contents = parser.parse()
        n_seg = sum(1 for _, _, _, ns, _ in parser._deferred_texts if ns)
        parser.apply_segmentation(contents, [["Body text with footnote."] for _ in range(n_seg)])
        parser.create_content_sections(contents)

        from bibr.paper_contents import CanonicalSection

        fn_sections = [s for s in contents.sections if s.section_type == CanonicalSection.FOOTNOTE]
        assert len(fn_sections) == 1
        # Footnote content present as a sentence
        fn_sentences = [s for s in contents.sentences if s.section_id == fn_sections[0].section_id]
        assert any("This is a footnote." in s.text for s in fn_sentences)
        # Xref linking footnote section back to the body sentence; xref_id is
        # the 1-based footnote ordinal (matches PDFParser, see PaperXref docs).
        xrefs = [x for x in contents.xrefs if x.xref_type == "foot"]
        assert len(xrefs) == 1
        assert xrefs[0].xref_id == 1


# ----- Run-level separators (w:br / w:tab / w:cr) -----


def _deferred_texts(build_fn) -> list[str]:
    """Paragraph text as the parser buffered it, before segmentation."""
    parser = DocxParser(_make_docx_bytes(build_fn))
    parser.parse()
    return [entry[0] for entry in parser._deferred_texts]


class TestRunSeparators:
    """python-docx writes plain strings, so no fixture ever emitted these."""

    def test_line_break_separates_words(self):
        def build(doc):
            para = doc.add_paragraph()
            run = para.add_run("Cognitive load and recall")
            run.add_break()
            run.add_text("Jane Smith")

        assert _deferred_texts(build) == ["Cognitive load and recall Jane Smith"]

    def test_tab_separates_words(self):
        def build(doc):
            para = doc.add_paragraph()
            run = para.add_run("Table 1")
            run.add_tab()
            run.add_text("Descriptive statistics")

        assert _deferred_texts(build) == ["Table 1 Descriptive statistics"]

    def test_a_shift_enter_title_block_stays_readable(self):
        def build(doc):
            para = doc.add_paragraph()
            run = para.add_run("Cognitive load and recall")
            for line in ("Jane Smith", "Department of Psychology", "jane.smith@example.edu"):
                run.add_break()
                run.add_text(line)

        assert _deferred_texts(build) == [
            "Cognitive load and recall Jane Smith Department of Psychology jane.smith@example.edu"
        ]

    def test_adjacent_runs_still_concatenate_untouched(self):
        """Word splits runs mid-word for formatting — no separator may appear."""

        def build(doc):
            para = doc.add_paragraph()
            para.add_run("Hyper")
            para.add_run("tension").bold = True

        assert _deferred_texts(build) == ["Hypertension"]

    def test_a_break_next_to_existing_whitespace_is_not_doubled(self):
        def build(doc):
            para = doc.add_paragraph()
            run = para.add_run("Methods ")
            run.add_break()
            run.add_text("We did things.")

        assert _deferred_texts(build) == ["Methods We did things."]


class TestNoteSeparators:
    def test_breaks_and_paragraphs_separate_endnote_text(self):
        from lxml import etree

        from bibr.input.docx_footnotes import load_endnotes

        w = _NS_W
        xml = f"""
        <w:endnotes xmlns:w="{w}">
          <w:endnote w:id="1">
            <w:p><w:r><w:t>Smith, J.</w:t><w:tab/><w:t>2020</w:t>
              <w:br/><w:t>Title of the work</w:t></w:r></w:p>
            <w:p><w:r><w:t>Journal of Testing.</w:t></w:r></w:p>
          </w:endnote>
        </w:endnotes>
        """

        class _FakePart:
            element = etree.fromstring(xml.strip())

        class _FakeDoc:
            class part:
                endnotes_part = _FakePart()

        assert load_endnotes(_FakeDoc()) == {
            "1": "Smith, J. 2020 Title of the work Journal of Testing."
        }
