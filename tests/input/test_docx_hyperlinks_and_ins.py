"""Tests for DOCX walker handling of ``w:hyperlink`` and ``w:ins`` wrappers.

The default walk recurses into ``w:r`` directly off the paragraph element.
``w:hyperlink`` (URL hyperlinks) and ``w:ins`` (accepted tracked-change
insertions) wrap their ``w:r`` children, so without explicit branches the
entire subtree is silently dropped from extracted text.
"""

import io

import pytest

pytest.importorskip("docx")  # python-docx


def _build_docx_with_hyperlink_and_insertion() -> bytes:
    """Build a DOCX whose paragraph contains:

    * a normal "hello " run,
    * a ``w:hyperlink`` wrapping a run pointing at https://example.org, and
    * a ``w:ins`` wrapping a run with inserted text.
    """
    from docx import Document
    from docx.oxml.ns import qn
    from lxml import etree

    doc = Document()
    p = doc.add_paragraph("hello ")
    rel_id = p.part.relate_to(
        "https://example.org",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    hyperlink = etree.SubElement(p._element, qn("w:hyperlink"))
    hyperlink.set(qn("r:id"), rel_id)
    r_in_link = etree.SubElement(hyperlink, qn("w:r"))
    t_in_link = etree.SubElement(r_in_link, qn("w:t"))
    t_in_link.text = "click here"

    p._element.append(
        etree.fromstring(
            f'<w:ins xmlns:w="{w_ns}" w:id="1" w:author="x" w:date="2024-01-01T00:00:00Z">'
            f"  <w:r><w:t> [inserted-tail]</w:t></w:r>"
            f"</w:ins>"
        )
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _parse_with_trivial_segmentation():
    """Parse the fixture DOCX, applying a trivial 1-segment-per-deferred segmentation."""
    from bibr.input.docx_native import DocxParser

    parser = DocxParser(_build_docx_with_hyperlink_and_insertion())
    contents = parser.parse()
    segments_per_entry = [
        [text] for text, _, _, needs_seg, _ in parser._deferred_texts if needs_seg
    ]
    parser.apply_segmentation(contents, segments_per_entry)
    return contents


def test_hyperlink_text_and_url_are_captured():
    contents = _parse_with_trivial_segmentation()

    all_text = " ".join(s.text for s in contents.sentences)
    assert "click here" in all_text, all_text
    assert any(link.url == "https://example.org" for link in contents.links)


def test_inserted_text_is_captured():
    contents = _parse_with_trivial_segmentation()

    all_text = " ".join(s.text for s in contents.sentences)
    assert "[inserted-tail]" in all_text, all_text


def _build_docx_with_url_as_display_text() -> bytes:
    """Build a DOCX with a hyperlink whose display text equals the URL string.

    Mirrors the common academic-paper DOI pattern, e.g. an inline link
    ``https://doi.org/10.1000/xyz`` shown verbatim.
    """
    from docx import Document
    from docx.oxml.ns import qn
    from lxml import etree

    url = "https://doi.org/10.1000/xyz"

    doc = Document()
    p = doc.add_paragraph("see ")
    rel_id = p.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )

    hyperlink = etree.SubElement(p._element, qn("w:hyperlink"))
    hyperlink.set(qn("r:id"), rel_id)
    r_in_link = etree.SubElement(hyperlink, qn("w:r"))
    t_in_link = etree.SubElement(r_in_link, qn("w:t"))
    t_in_link.text = url

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_hyperlink_with_url_as_display_text_does_not_double_link():
    """When hyperlink display text IS the URL string, only one PaperURLLink
    is emitted (not one from the hyperlink branch + one from the regex pass)."""
    from bibr.input.docx_native import DocxParser

    url = "https://doi.org/10.1000/xyz"
    parser = DocxParser(_build_docx_with_url_as_display_text())
    contents = parser.parse()
    segments_per_entry = [
        [text] for text, _, _, needs_seg, _ in parser._deferred_texts if needs_seg
    ]
    parser.apply_segmentation(contents, segments_per_entry)

    matching = [link for link in contents.links if link.url == url]
    assert len(matching) == 1, [(link.url, link.link_text, link.text_id) for link in contents.links]
