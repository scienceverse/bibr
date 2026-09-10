"""M14: DOCX content the native parser walked straight past.

Endnotes (never loaded at all — an endnote-style bibliography, the humanities
convention, vanished entirely), ``w:sdt`` content controls (Word wraps cover
pages, abstract boxes and generated bibliographies in these), and text boxes
inside a drawing.
"""

import io
import zipfile

import pytest

from bibr.input.docx_footnotes import load_endnotes, load_footnotes

pytest.importorskip("docx")

_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_R = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'

_CONTENT_TYPES = """<?xml version="1.0"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/word/endnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml"/>
  <Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>
</Types>"""

_ROOT_RELS = """<?xml version="1.0"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""

_DOC_RELS = """<?xml version="1.0"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId10" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes" Target="endnotes.xml"/>
  <Relationship Id="rId11" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>
</Relationships>"""

_ENDNOTES = f"""<?xml version="1.0"?>
<w:endnotes {_W}>
  <w:endnote w:type="separator" w:id="-1"><w:p><w:r><w:t>sep</w:t></w:r></w:p></w:endnote>
  <w:endnote w:id="1"><w:p><w:r><w:t>Foucault, M. (1975). Surveiller et punir.</w:t></w:r></w:p></w:endnote>
  <w:endnote w:id="2"><w:p><w:r><w:t>Bourdieu, P. (1979). La distinction.</w:t></w:r></w:p></w:endnote>
</w:endnotes>"""

_FOOTNOTES = f"""<?xml version="1.0"?>
<w:footnotes {_W}>
  <w:footnote w:id="1"><w:p><w:r><w:t>A footnote, not an endnote.</w:t></w:r></w:p></w:footnote>
</w:footnotes>"""

_DOCUMENT = f"""<?xml version="1.0"?>
<w:document {_W} {_R}>
  <w:body>
    <w:p><w:r><w:t>Discipline emerged as a technique.</w:t>
      <w:endnoteReference w:id="1"/></w:r></w:p>
    <w:sdt>
      <w:sdtPr><w:alias w:val="Bibliography"/></w:sdtPr>
      <w:sdtContent>
        <w:p><w:r><w:t>Text inside a content control.</w:t></w:r></w:p>
      </w:sdtContent>
    </w:sdt>
    <w:p><w:r><w:t>Taste is a marker of class.</w:t>
      <w:endnoteReference w:id="2"/></w:r></w:p>
  </w:body>
</w:document>"""


def _docx_bytes(document=_DOCUMENT):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", _CONTENT_TYPES)
        z.writestr("_rels/.rels", _ROOT_RELS)
        z.writestr("word/document.xml", document)
        z.writestr("word/_rels/document.xml.rels", _DOC_RELS)
        z.writestr("word/endnotes.xml", _ENDNOTES)
        z.writestr("word/footnotes.xml", _FOOTNOTES)
    return buf.getvalue()


def _parse(document=_DOCUMENT):
    from bibr.input.docx_native import DocxParser

    parser = DocxParser(_docx_bytes(document))
    contents = parser.parse()
    parser.apply_segmentation(
        contents, [[e.text] for e in parser.assembler.entries if e.needs_segmentation]
    )
    parser.create_content_sections(contents)
    return contents


class TestEndnoteLoading:
    def test_endnotes_are_read_from_their_own_part(self):
        import docx

        doc = docx.Document(io.BytesIO(_docx_bytes()))
        notes = load_endnotes(doc)

        assert notes == {
            "1": "Foucault, M. (1975). Surveiller et punir.",
            "2": "Bourdieu, P. (1979). La distinction.",
        }

    def test_separators_are_skipped(self):
        import docx

        doc = docx.Document(io.BytesIO(_docx_bytes()))
        assert "-1" not in load_endnotes(doc)

    def test_footnotes_still_load_from_their_own_id_space(self):
        import docx

        doc = docx.Document(io.BytesIO(_docx_bytes()))
        # Both parts have an id "1"; they must not be merged.
        assert load_footnotes(doc)["1"] == "A footnote, not an endnote."
        assert load_endnotes(doc)["1"].startswith("Foucault")


class TestEndnotesReachTheDocument:
    def test_endnote_text_becomes_sentences(self):
        contents = _parse()
        text = " ".join(s.text for s in contents.sentences)

        assert "Surveiller et punir" in text
        assert "La distinction" in text

    def test_endnote_sections_are_labelled_as_endnotes(self):
        contents = _parse()
        headers = [s.header for s in contents.sections]

        assert "Endnote 1" in headers
        assert "Endnote 2" in headers

    def test_endnotes_are_numbered_within_their_own_kind(self):
        contents = _parse()
        foot_xrefs = [x for x in contents.xrefs if x.xref_type == "foot"]

        assert [x.contents for x in foot_xrefs] == ["1", "2"]


class TestContentControls:
    def test_paragraphs_inside_an_sdt_are_parsed(self):
        contents = _parse()
        text = " ".join(s.text for s in contents.sentences)

        assert "Text inside a content control." in text


class TestTextBoxes:
    def test_text_inside_a_drawing_text_box_is_kept(self):
        document = f"""<?xml version="1.0"?>
<w:document {_W} {_R}>
  <w:body>
    <w:p><w:r>
      <w:t>Body sentence.</w:t>
      <w:drawing>
        <w:txbxContent>
          <w:p><w:r><w:t>Boxed pull quote.</w:t></w:r></w:p>
        </w:txbxContent>
      </w:drawing>
    </w:r></w:p>
  </w:body>
</w:document>"""
        contents = _parse(document)
        text = " ".join(s.text for s in contents.sentences)

        assert "Boxed pull quote." in text
