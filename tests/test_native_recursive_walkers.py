"""Regression tests for audit action 5 (native recursive walkers).

Each finding below failed on base (see FINDINGS.md): the test pins the fixed
behaviour, and the guard next to it pins the boundary the fix must not cross.
"""

from __future__ import annotations

import base64
import io

import pytest

pytest.importorskip("docx")

from docx import Document  # noqa: E402

from bibr.input.docx_native import DocxParser  # noqa: E402
from bibr.input.html_native import HtmlParser  # noqa: E402

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _html_texts(data: bytes) -> list[str]:
    parser = HtmlParser(data)
    parser.parse()
    return [entry.text for entry in parser.assembler.entries]


def _docx_texts(docx_bytes: bytes) -> tuple[DocxParser, list[str]]:
    parser = DocxParser(docx_bytes)
    parser.parse()
    return parser, [entry.text for entry in parser.assembler.entries]


def _save(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _raw_paragraph(doc: Document, inner_xml: str):
    """Append a paragraph whose content is the given raw w: XML."""
    from lxml import etree

    para = doc.add_paragraph()
    for run in list(para._element.findall(f"{{{_W}}}r")):
        para._element.remove(run)
    frag = etree.fromstring(f'<w:p xmlns:w="{_W}" xmlns:m="{_M}">{inner_xml}</w:p>')
    for child in list(frag):
        para._element.append(child)
    return para


class TestDivDirectText:
    def test_div_paragraphs_are_kept_with_inline_children(self):
        texts = _html_texts(
            b"""<html><body><article><h2>Introduction</h2>"""
            b"""<div class="para">Direct div text is common. <a href="#bib1">[1]</a>.</div>"""
            b"""<p>Paragraph kept.</p></article></body></html>"""
        )
        assert texts[0] == "Direct div text is common. [1]."
        assert "Paragraph kept." in texts

    def test_bare_section_text_is_kept_in_order(self):
        texts = _html_texts(
            b"""<html><body><section><h1>Chapter</h1>Bare section text here."""
            b"""<p>Kept p text.</p></section></body></html>"""
        )
        assert texts == ["Bare section text here.", "Kept p text."]

    def test_nested_div_paragraphs_stay_separate_entries(self):
        texts = _html_texts(
            b"""<html><body><article><h2>Intro</h2>"""
            b"""<div><div class="para">First div para.</div>"""
            b"""<div class="para">Second div para.</div></div>"""
            b"""</article></body></html>"""
        )
        assert texts == ["First div para.", "Second div para."]

    def test_div_with_block_children_recurses(self):
        texts = _html_texts(
            b"""<html><body><article><h2>Intro</h2>"""
            b"""<div><p>One.</p><p>Two.</p></div></article></body></html>"""
        )
        assert texts == ["One.", "Two."]

    def test_noise_tags_still_dropped(self):
        texts = _html_texts(
            b"""<html><body><article><h2>Intro</h2>"""
            b"""<div>Kept.<script>var x = 1;</script><style>.a{}</style></div>"""
            b"""</article></body></html>"""
        )
        assert texts == ["Kept."]

    def test_inline_wrapper_around_blocks_recurses(self):
        # An unclosed <b> nests whole sections inside an inline element;
        # the block structure wins instead of flattening into body text.
        parser = HtmlParser(
            b"""<html><body><article><h2>Intro</h2><p>Body.</p>"""
            b"""<b><section><h2>Footer head</h2><p>Footer text.</p></section></b>"""
            b"""</article></body></html>"""
        )
        contents = parser.parse()
        assert [s.header for s in contents.sections] == ["Root", "Intro", "Footer head"]
        assert [e.text for e in parser.assembler.entries] == ["Body.", "Footer text."]


class TestHtmlCharset:
    _BODY = "<html><body><article><h2>%s</h2><p>%s</p></article></body></html>"

    def test_utf8_without_meta_charset(self):
        data = (self._BODY % ("Über Größe", "Café naïve résumé")).encode("utf-8")
        assert _html_texts(data) == ["Café naïve résumé"]

    def test_xml_declaration_encoding_is_honoured(self):
        data = (
            '<?xml version="1.0" encoding="UTF-8"?>' + self._BODY % ("Über Größe", "Café")
        ).encode("utf-8")
        parser = HtmlParser(data)
        contents = parser.parse()
        assert contents.detected_title == "Über Größe"
        assert _html_texts(data) == ["Café"]

    def test_declared_windows_1252_still_decodes(self):
        data = (
            '<html><head><meta charset="windows-1252"></head>'
            "<body><article><h2>Intro</h2><p>Caf\xe9</p></article></body></html>"
        ).encode("latin-1")
        assert _html_texts(data) == ["Café"]


class TestFigureWrappedTable:
    def test_figure_wrapping_a_table_becomes_a_table(self):
        parser = HtmlParser(
            b"""<html><body><article><h2>Results</h2>"""
            b"""<figure class="table"><figcaption>Table 1. Demographics</figcaption>"""
            b"""<table><tr><th>A</th></tr><tr><td>1</td></tr></table></figure>"""
            b"""<p>Body.</p></article></body></html>"""
        )
        contents = parser.parse()
        assert contents.figures == []
        (table,) = contents.tables
        assert table.caption == "Table 1. Demographics"
        assert table.label == "1"

    def test_figure_with_an_image_and_a_layout_table_stays_a_figure(self):
        parser = HtmlParser(
            b"""<html><body><article><h2>Results</h2>"""
            b"""<figure><img src="a.png"><figcaption>Figure 1. Main.</figcaption>"""
            b"""<table><tr><td>x</td></tr></table></figure>"""
            b"""</article></body></html>"""
        )
        contents = parser.parse()
        assert [fig.caption for fig in contents.figures] == ["Figure 1. Main."]
        assert contents.tables == []


class TestHtmlAuthorMetadata:
    _HEAD = "<html><head>%s</head><body><article><h2>Intro</h2><p>Body.</p></article></body></html>"

    def _meta(self, head: str):
        parser = HtmlParser((self._HEAD % head).encode())
        parser.parse()
        return parser._metadata

    def test_first_source_wins_and_institutions_pair_up(self):
        meta = self._meta(
            '<meta name="citation_author" content="Jane Smith">'
            '<meta name="citation_author_institution" content="Utrecht University">'
            '<meta name="citation_author" content="John Doe">'
            '<meta name="dc.creator" content="Jane Smith">'
            '<meta name="dc.creator" content="John Doe">'
        )
        assert [(a.given, a.family, a.affiliation) for a in meta.authors] == [
            ("Jane", "Smith", "Utrecht University"),
            ("John", "Doe", ""),
        ]

    def test_institutions_attach_to_the_preceding_author(self):
        # Highwire order: each institution follows its own author, and one
        # author can carry several. Positional zipping gave Doe "KNAW"
        # (Smith's second institution) and Lee's email/ORCID to Smith.
        meta = self._meta(
            '<meta name="citation_author" content="Jane Smith">'
            '<meta name="citation_author_institution" content="Utrecht University">'
            '<meta name="citation_author_institution" content="KNAW">'
            '<meta name="citation_author" content="John Doe">'
            '<meta name="citation_author_institution" content="Leiden University">'
            '<meta name="citation_author" content="Ann Lee">'
            '<meta name="citation_author_institution" content="Oxford">'
            '<meta name="citation_author_email" content="ann.lee@example.org">'
            '<meta name="citation_author_orcid" content="0000-0002-1825-0097">'
        )
        assert [(a.given, a.family, a.affiliation) for a in meta.authors] == [
            ("Jane", "Smith", "Utrecht University; KNAW"),
            ("John", "Doe", "Leiden University"),
            ("Ann", "Lee", "Oxford"),
        ]
        assert [a.email for a in meta.authors] == [None, None, "ann.lee@example.org"]
        assert [a.orcid for a in meta.authors] == [
            None,
            None,
            "https://orcid.org/0000-0002-1825-0097",
        ]

    def test_non_doi_identifier_is_not_a_doi(self):
        meta = self._meta(
            '<meta name="citation_title" content="T">'
            '<meta name="citation_author" content="Jane Smith">'
            '<meta name="dc.identifier"'
            ' content="https://journal.example.org/article/view/123">'
        )
        assert meta.doi == ""

    def test_bare_doi_identifier_normalizes(self):
        meta = self._meta(
            '<meta name="citation_title" content="T">'
            '<meta name="citation_author" content="Jane Smith">'
            '<meta name="dc.identifier" content="https://doi.org/10.1234/abc.1">'
        )
        assert meta.doi == "10.1234/abc.1"

    def test_generic_author_list_splits_but_family_given_stays(self):
        meta = self._meta('<meta name="author" content="Jane Smith, John Doe">')
        assert [(a.given, a.family) for a in meta.authors] == [
            ("Jane", "Smith"),
            ("John", "Doe"),
        ]
        single = self._meta('<meta name="author" content="Smith, Jane">')
        assert [(a.given, a.family) for a in single.authors] == [("Jane", "Smith")]

    def test_citation_orcid_pairs_up(self):
        meta = self._meta(
            '<meta name="citation_author" content="Jane Smith">'
            '<meta name="citation_author_orcid" content="0000-0002-1825-0097">'
        )
        assert meta.authors[0].orcid == "https://orcid.org/0000-0002-1825-0097"


class TestPreparsedGating:
    _SEO = (
        b"""<html><head><title>Sleep and memory | Journal of Sleep | Example Press</title>"""
        b"""<meta name="description" content="Read the latest research from Example Press.">"""
        b"""<meta name="author" content="Jane Smith, John Doe"></head>"""
        b"""<body><article><h1>Sleep and memory</h1><h2>Abstract</h2>"""
        b"""<p>We tested 80 people.</p></article></body></html>"""
    )
    _CITED = (
        b"""<html><head><meta name="citation_title" content="Cited Paper">"""
        b"""<meta name="citation_author" content="Jane Smith">"""
        b"""<meta name="citation_abstract" content="Author-written abstract."></head>"""
        b"""<body><article><h1>Cited Paper</h1><h2>Abstract</h2>"""
        b"""<p>We tested 80 people.</p></article></body></html>"""
    )

    def test_seo_only_page_is_marked_untrusted(self):
        # The record stays (a no-LLM run keeps its language and licence);
        # the flag tells an LLM run to extract the front matter instead.
        contents = HtmlParser(self._SEO).parse()
        assert contents.preparsed_metadata is not None
        assert contents.preparsed_metadata_trusted is False

    def test_seo_description_is_not_the_abstract(self):
        parser = HtmlParser(self._SEO)
        parser.parse()
        assert parser._metadata.abstract == ""

    def test_site_suffix_strips_when_h1_matches(self):
        parser = HtmlParser(self._SEO)
        parser.parse()
        assert parser._metadata.title == "Sleep and memory"

    def test_hyphenated_title_without_suffix_spacing_is_untouched(self):
        parser = HtmlParser(
            b"""<html><head><title>Sleep-memory study</title></head>"""
            b"""<body><article><h1>Sleep</h1><p>Body.</p></article></body></html>"""
        )
        parser.parse()
        assert parser._metadata.title == "Sleep-memory study"

    def test_citation_front_matter_still_preparsed(self):
        contents = HtmlParser(self._CITED).parse()
        assert contents.preparsed_metadata is not None
        assert contents.preparsed_metadata_trusted is True
        assert contents.preparsed_metadata.abstract == "Author-written abstract."


class TestDocxInlineWalker:
    def test_inline_sdt_citation_is_kept(self):
        doc = Document()
        _raw_paragraph(
            doc,
            "<w:r><w:t>As shown previously </w:t></w:r>"
            '<w:sdt><w:sdtPr><w:alias w:val="citation"/></w:sdtPr>'
            "<w:sdtContent><w:r><w:t>(Smith et al., 2020)</w:t></w:r></w:sdtContent></w:sdt>"
            "<w:r><w:t>, effects were large.</w:t></w:r>",
        )
        _, texts = _docx_texts(_save(doc))
        assert texts == ["As shown previously (Smith et al., 2020), effects were large."]

    def test_smart_tag_fldsimple_sym_and_hyphen(self):
        doc = Document()
        _raw_paragraph(
            doc,
            "<w:r><w:t>Data collected in </w:t></w:r>"
            "<w:smartTag><w:r><w:t>Amsterdam</w:t></w:r></w:smartTag>"
            "<w:r><w:t> in 2019. See </w:t></w:r>"
            '<w:fldSimple w:instr="REF _Ref1"><w:r><w:t>Table 2</w:t></w:r></w:fldSimple>'
            "<w:r><w:t> for details. Level </w:t></w:r>"
            '<w:r><w:sym w:font="Symbol" w:char="F061"/></w:r>'
            "<w:r><w:t> = .05 for COVID</w:t></w:r>"
            "<w:r><w:noBreakHyphen/><w:t>19 cases.</w:t></w:r>",
        )
        _, texts = _docx_texts(_save(doc))
        assert texts == [
            "Data collected in Amsterdam in 2019. "
            "See Table 2 for details. Level \u03b1 = .05 for COVID-19 cases."
        ]

    def test_tracked_move_destination_is_kept(self):
        doc = Document()
        _raw_paragraph(
            doc,
            "<w:moveTo><w:r><w:t>Moved sentence that should remain.</w:t></w:r></w:moveTo>",
        )
        _, texts = _docx_texts(_save(doc))
        assert texts == ["Moved sentence that should remain."]

    def test_deletions_move_sources_and_field_code_stay_skipped(self):
        doc = Document()
        _raw_paragraph(
            doc,
            "<w:r><w:t>Kept</w:t></w:r>"
            "<w:del><w:r><w:delText>gone</w:delText></w:r></w:del>"
            "<w:moveFrom><w:r><w:t>gone</w:t></w:r></w:moveFrom>"
            "<w:r><w:instrText>PAGE</w:instrText></w:r>"
            "<w:r><w:t> here.</w:t></w:r>",
        )
        _, texts = _docx_texts(_save(doc))
        assert texts == ["Kept here."]

    def test_complex_field_result_runs_are_kept(self):
        doc = Document()
        _raw_paragraph(
            doc,
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            "<w:r><w:instrText> CITATION Smi20 </w:instrText></w:r>"
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            "<w:r><w:t>(Smith, 2020)</w:t></w:r>"
            '<w:r><w:fldChar w:fldCharType="end"/></w:r>',
        )
        _, texts = _docx_texts(_save(doc))
        assert texts == ["(Smith, 2020)"]


class TestDocxHeadingFootnotes:
    def _footnote_doc(self):
        import zipfile

        from lxml import etree

        doc = Document()
        title = doc.add_paragraph("Cognitive load and recall")
        title.style = doc.styles["Title"]
        doc.add_heading("Intro heading", level=1)
        doc.add_paragraph("Body text.")
        data = _save(doc)

        parts = {}
        with zipfile.ZipFile(io.BytesIO(data)) as zin:
            for name in zin.namelist():
                parts[name] = zin.read(name)
        root = etree.fromstring(parts["word/document.xml"])
        ns = {"w": _W}
        paras = root.findall(".//w:p", ns)
        for index, note_id in enumerate(["1", "2"]):
            run = paras[index].find("w:r", ns)
            ref = etree.SubElement(run, f"{{{_W}}}footnoteReference")
            ref.set(f"{{{_W}}}id", note_id)
        parts["word/document.xml"] = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        parts["word/footnotes.xml"] = (
            f'<w:footnotes xmlns:w="{_W}">'
            '<w:footnote w:id="1"><w:p><w:r><w:t>Author note funded.</w:t></w:r></w:p></w:footnote>'
            '<w:footnote w:id="2"><w:p><w:r><w:t>Preregistered note.</w:t></w:r></w:p></w:footnote>'
            "</w:footnotes>"
        ).encode()
        cts = parts["[Content_Types].xml"].decode()
        if "footnotes" not in cts:
            cts = cts.replace(
                "</Types>",
                '<Override PartName="/word/footnotes.xml" '
                'ContentType="application/vnd.openxmlformats-officedocument'
                '.wordprocessingml.footnotes+xml"/></Types>',
            )
            parts["[Content_Types].xml"] = cts.encode()
        rels = parts["word/_rels/document.xml.rels"].decode()
        if "footnotes" not in rels:
            rels = rels.replace(
                "</Relationships>",
                '<Relationship Id="rFoot1" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006'
                '/relationships/footnotes" Target="footnotes.xml"/></Relationships>',
            )
            parts["word/_rels/document.xml.rels"] = rels.encode()
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for name, blob in parts.items():
                zout.writestr(name, blob)
        return buf.getvalue()

    def test_heading_anchored_footnotes_are_emitted(self):
        parser = DocxParser(self._footnote_doc())
        contents = parser.parse()
        parser.apply_segmentation(
            contents,
            [[entry.text] for entry in parser.assembler.entries if entry.needs_segmentation],
        )
        parser.create_content_sections(contents)
        headers = [section.header for section in contents.sections]
        assert "Footnote 1" in headers
        assert "Footnote 2" in headers
        text = " ".join(sentence.text for sentence in contents.sentences)
        assert "Author note funded." in text
        assert "Preregistered note." in text


class TestDocxCustomHeadings:
    def test_style_based_on_heading_is_a_heading(self):
        doc = Document()
        based = doc.styles.add_style("APA Heading 1", 1)
        based.base_style = doc.styles["Heading 1"]
        doc.add_paragraph("Method", style="APA Heading 1")
        doc.add_paragraph("We recruited 80 participants.")
        doc.add_paragraph("Results", style="APA Heading 1")
        doc.add_paragraph("Effects were large.")
        contents = DocxParser(_save(doc)).parse()
        assert [s.header for s in contents.sections] == ["Root", "Method", "Results"]

    def test_outline_level_marks_a_heading(self):
        from lxml import etree

        doc = Document()
        para = doc.add_paragraph("Discussion")
        ppr = para._element.get_or_add_pPr()
        lvl = etree.SubElement(ppr, f"{{{_W}}}outlineLvl")
        lvl.set(f"{{{_W}}}val", "1")
        doc.add_paragraph("Body.")
        contents = DocxParser(_save(doc)).parse()
        assert [s.header for s in contents.sections] == ["Root", "Discussion"]

    def test_normal_style_stays_body_text(self):
        doc = Document()
        doc.add_paragraph("Just a sentence.")
        contents = DocxParser(_save(doc)).parse()
        assert [s.header for s in contents.sections] == ["Root"]


class TestDocxDetectedTitle:
    def test_hand_formatted_title_does_not_promote_introduction(self):
        doc = Document()
        lead = doc.add_paragraph("My Hand-Formatted Title")
        for run in lead.runs:
            run.bold = True
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph("Intro body.")
        doc.add_heading("Method", level=1)
        doc.add_paragraph("Method body.")
        contents = DocxParser(_save(doc)).parse()
        assert contents.detected_title != "Introduction"

    def test_title_style_still_detected(self):
        doc = Document()
        title = doc.add_paragraph("Real Title")
        title.style = doc.styles["Title"]
        doc.add_heading("Introduction", level=1)
        contents = DocxParser(_save(doc)).parse()
        assert contents.detected_title == "Real Title"

    def test_first_heading_one_that_is_not_a_section_stays_title(self):
        doc = Document()
        doc.add_heading("Paper About X", level=1)
        doc.add_paragraph("Body text.")
        contents = DocxParser(_save(doc)).parse()
        assert contents.detected_title == "Paper About X"


class TestDocxFigureCaptions:
    def _picture_doc(self, build):
        doc = Document()
        build(doc)
        return _save(doc)

    def _add_pic(self, doc):
        from docx.shared import Pt

        doc.add_paragraph().add_run().add_picture(io.BytesIO(_PNG), width=Pt(10))

    def test_icon_does_not_shift_later_captions(self):
        def build(doc):
            doc.add_paragraph("Intro text.")
            self._add_pic(doc)
            self._add_pic(doc)
            doc.add_paragraph("Figure 1. Main result.", style="Caption")
            self._add_pic(doc)
            doc.add_paragraph("Figure 2. Second result.", style="Caption")

        contents = DocxParser(self._picture_doc(build)).parse()
        assert [(fig.caption, fig.label) for fig in contents.figures] == [
            (None, None),
            ("Figure 1. Main result.", "1"),
            ("Figure 2. Second result.", "2"),
        ]

    def test_caption_above_the_figure_pairs(self):
        def build(doc):
            doc.add_paragraph("Figure 1. Above.", style="Caption")
            self._add_pic(doc)

        contents = DocxParser(self._picture_doc(build)).parse()
        assert [(fig.caption, fig.label) for fig in contents.figures] == [("Figure 1. Above.", "1")]

    def test_body_text_between_picture_and_caption_breaks_pairing(self):
        def build(doc):
            self._add_pic(doc)
            doc.add_paragraph("Some body text.")
            doc.add_paragraph("Figure 1. Orphan.", style="Caption")

        contents = DocxParser(self._picture_doc(build)).parse()
        assert [fig.caption for fig in contents.figures] == [None]


class TestDcFrontMatterGating:
    """Dublin Core / editorial front matter must survive as preparsed metadata.

    Gating on citation_title plus citation_author discarded every eLife page
    (dc.title plus a DOI, no citation_* tags): the export lost the article
    DOI to a component DOI picked from body text, plus the published date,
    publisher, license and language. Fails on the gated code (untrusted).
    """

    _DC = (
        b"""<html lang="en"><head>"""
        b"""<meta name="dc.title" content="A bacterial sulfonolipid triggers development">"""
        b"""<meta name="dc.identifier" content="doi:10.7554/eLife.00013">"""
        b"""<meta name="dc.date" content="2012-10-15">"""
        b"""<meta name="dc.publisher" content="eLife Sciences Publications Limited">"""
        b"""<meta name="dc.rights" content="CC-BY license text">"""
        b"""<meta name="dc.language" content="en"></head>"""
        b"""<body><main><h2>Introduction</h2><p>Body.</p></main></body></html>"""
    )

    def test_dc_only_page_keeps_its_doi_and_front_matter(self):
        contents = HtmlParser(self._DC).parse()
        assert contents.preparsed_metadata_trusted is True
        meta = contents.preparsed_metadata
        assert meta is not None
        assert meta.doi == "10.7554/eLife.00013"
        assert meta.title == "A bacterial sulfonolipid triggers development"
        assert meta.published == "2012-10-15"
        assert meta.publisher == "eLife Sciences Publications Limited"
        assert meta.license == "CC-BY license text"
        assert meta.language == "en"

    def test_title_and_doi_without_any_author_keeps_preparsed(self):
        # Editorials/errata and ePub OPF records (title plus DOI, no
        # creator) carry structured identity without authors.
        html = (
            b"""<html><head><meta name="citation_title" content="Editorial">"""
            b"""<meta name="citation_doi" content="10.1234/ed.2020"></head>"""
            b"""<body><article><h2>Intro</h2><p>Text.</p></article></body></html>"""
        )
        contents = HtmlParser(html).parse()
        assert contents.preparsed_metadata_trusted is True
        meta = contents.preparsed_metadata
        assert meta is not None
        assert (meta.title, meta.doi) == ("Editorial", "10.1234/ed.2020")

    def test_dc_title_with_dc_creator_reaches_the_export(self):
        # The dc.creator branch is only real if the gate lets it through.
        html = (
            b"""<html><head><meta name="dc.title" content="DC Paper">"""
            b"""<meta name="dc.creator" content="Jane Smith">"""
            b"""<meta name="dc.creator" content="John Doe"></head>"""
            b"""<body><article><h2>Intro</h2><p>Text.</p></article></body></html>"""
        )
        contents = HtmlParser(html).parse()
        assert contents.preparsed_metadata_trusted is True
        meta = contents.preparsed_metadata
        assert meta is not None
        assert [(a.given, a.family) for a in meta.authors] == [
            ("Jane", "Smith"),
            ("John", "Doe"),
        ]

    def test_three_generic_authors_split_and_reach_the_export(self):
        # "Jane Smith, John Doe, Ann Lee" split into two on the first comma
        # only, parsing as family "John Doe", given "Ann Lee".
        html = (
            b"""<html><head><meta name="dc.title" content="Generic Paper">"""
            b"""<meta name="author" content="Jane Smith, John Doe, Ann Lee"></head>"""
            b"""<body><article><h2>Intro</h2><p>Text.</p></article></body></html>"""
        )
        contents = HtmlParser(html).parse()
        assert contents.preparsed_metadata_trusted is True
        meta = contents.preparsed_metadata
        assert meta is not None
        assert [(a.given, a.family) for a in meta.authors] == [
            ("Jane", "Smith"),
            ("John", "Doe"),
            ("Ann", "Lee"),
        ]


class TestLegacyCharsetDecoding:
    """WHATWG label resolution for declared and undeclared legacy pages."""

    def test_iso_8859_1_label_reads_cp1252_punctuation(self):
        # WHATWG maps iso-8859-1 to windows-1252; Python's codec alone emits
        # C1 controls for the cp1252 smart quotes legacy pages rely on.
        for meta in (
            b'<meta charset="iso-8859-1">',
            b'<meta http-equiv="Content-Type" content="text/html; charset=iso-8859-1">',
        ):
            data = (
                b"<html><head>"
                + meta
                + b"</head><body><article><h2>I</h2><p>\x93Quoted\x94 \x96 dash and caf\xe9.</p></article></body></html>"
            )
            assert _html_texts(data) == ["\u201cQuoted\u201d \u2013 dash and caf\u00e9."]

    def test_undeclared_cp1252_reads_as_windows_1252(self):
        data = (
            "<html><body><article><h2>I</h2><p>Caf\xe9 na\xefve r\xe9sum\xe9.</p>"
            "</article></body></html>"
        ).encode("latin-1")
        assert _html_texts(data) == ["Caf\u00e9 na\u00efve r\u00e9sum\u00e9."]

    def test_utf16_label_on_8bit_bytes_does_not_decode_as_cjk(self):
        data = (
            b'<html><head><meta charset="utf-16"></head>'
            b"<body><article><h2>I</h2><p>caf\xe9 ok.</p></article></body></html>"
        )
        assert _html_texts(data) == ["caf\u00e9 ok."]

    def test_xml_declaration_label_uses_whatwg_mapping(self):
        # An XML-declared latin1 label with cp1252 bytes must come through
        # windows-1252, not Python's latin1 codec (C1 controls).
        body = "“Quoted” – dash.".encode("cp1252")
        raw = (
            b'<?xml version="1.0" encoding="latin1"?><html><body><main>'
            b"<section><h2>Note</h2><p>" + body + b"</p></section></main></body></html>"
        )
        assert b"\x93" in raw
        assert _html_texts(raw) == ["“Quoted” – dash."]

    def test_non_utf8_xml_declaration_encoding_is_honoured(self):
        # iso-8859-2 bytes that are invalid UTF-8 must come through the
        # XML-declared codec, not the replacement fallback.
        template = '<?xml version="1.0" encoding="iso-8859-2"?><html><body><article><h2>Miodowa</h2><p>%s</p></article></body></html>'
        data = (template % "Za\u017c\u00f3\u0142\u0107.").encode("iso-8859-2")
        assert _html_texts(data) == ["Za\u017c\u00f3\u0142\u0107."]


class TestPageChromeIsNoise:
    """Screen-reader-only markup, buttons and link-only containers are chrome."""

    _CHROME = (
        b"""<html><body><main>"""
        b"""<header><h2>Materials and methods</h2>"""
        b"""<a href="https://bio-protocol.example/s4-1">Request a detailed protocol</a></header>"""
        b"""<section><h3>Odor delivery</h3><p>We built a system.</p>"""
        b"""<span class="doi doi--article-section">"""
        b"""<a href="https://doi.org/10.7554/eLife.06651.001">https://doi.org/10.7554/eLife.06651.001</a>"""
        b"""</span></section>"""
        b"""<div><span class="visuallyhidden">Download asset</span>"""
        b"""<span class="visuallyhidden">Open asset</span></div>"""
        b"""<button type="button">Copy to clipboard</button>"""
        b"""<div aria-hidden="true">Figure 1 with 0 supplements</div>"""
        b"""</main></body></html>"""
    )

    def test_header_links_and_hidden_chrome_leave_no_text(self):
        texts = _html_texts(self._CHROME)
        assert texts == ["We built a system."]
        for dropped in (
            "Request a detailed protocol",
            "Download asset",
            "Open asset",
            "Copy to clipboard",
            "https://doi.org/10.7554/eLife.06651.001",
            "Figure 1 with 0 supplements",
        ):
            assert dropped not in " ".join(texts)

    def test_paragraph_with_a_link_and_words_is_kept(self):
        # Guard: the link-only rule must not eat prose that merely contains
        # a link.
        texts = _html_texts(
            b"""<html><body><main><section><h2>S</h2>"""
            b"""<div>See the <a href="https://example.org/data">dataset</a> for details.</div>"""
            b"""</section></main></body></html>"""
        )
        assert texts == ["See the dataset for details."]

    def test_cite_formatter_lists_are_not_reference_lists(self):
        # A "cite this article" block reuses reference__* BEM classes for a
        # single formatted citation; it must not open a References section,
        # emit author-name reference strings, or leave its parts in the body.
        parser = HtmlParser(
            b"""<html><body><main><section><h2>Download links</h2><p>Prose.</p>"""
            b"""<div class="reference"><ol class="reference__authors_list">"""
            b"""<li>Anmo J Kim</li><li>Aurel A Lazar</li></ol>"""
            b"""<span class="reference__authors_list_suffix">(2015)</span>"""
            b"""<div class="reference__title">Some title</div>"""
            b"""<div class="reference__origin"><i>eLife</i> <b>4</b>:e06651.</div></div>"""
            b"""</section></main></body></html>"""
        )
        contents = parser.parse()
        assert [s.header for s in contents.sections] == ["Root", "Download links"]
        assert (contents.native_ref_strings or []) == []
        assert [e.text for e in parser.assembler.entries] == ["Prose."]

    def test_real_reference_list_class_still_counts(self):
        # Guard: the BEM exclusion must keep genuine bibliography lists.
        parser = HtmlParser(
            b"""<html><body><main><section><h2>Body</h2>"""
            b"""<ol class="reference-list"><li>Smith J (2020) A title. J Psych 1:2.</li>"""
            b"""<li>Lee A (2019) Other. Nature 3:4.</li></ol>"""
            b"""</section></main></body></html>"""
        )
        contents = parser.parse()
        assert contents.native_ref_strings == [
            "Smith J (2020) A title. J Psych 1:2.",
            "Lee A (2019) Other. Nature 3:4.",
        ]


class TestInlineWhitespaceAndReferences:
    def test_whitespace_between_inline_siblings_is_kept(self):
        texts = _html_texts(
            b"""<html><body><main><section><h2>Intro</h2>"""
            b"""Bare <b>bold</b> words and <span>Smith J,</span> <span>Doe K</span> here."""
            b"""<p>Next.</p></section></main></body></html>"""
        )
        assert texts[0] == "Bare bold words and Smith J, Doe K here."

    def test_structured_div_reference_is_one_reference(self):
        parser = HtmlParser(
            b"""<html><body><main><section><h2>Body</h2><p>Prose.</p></section>"""
            b"""<section><h2>References</h2>"""
            b"""<div class="ref"><span>Smith J, Doe K</span> <span>(2020)</span>"""
            b"""<div>A title.</div><div>J Psych 1:2.</div></div>"""
            b"""<div class="ref"><span>Lee A</span> <span>(2019)</span>"""
            b"""<div>Other.</div><div>Nature 3:4.</div></div>"""
            b"""</section></main></body></html>"""
        )
        contents = parser.parse()
        assert contents.native_ref_strings == [
            "Smith J, Doe K (2020) A title. J Psych 1:2.",
            "Lee A (2019) Other. Nature 3:4.",
        ]

    def test_flat_div_references_stay_separate(self):
        # Guard: a wrapper holding only nested divs still recurses into one
        # reference per child.
        parser = HtmlParser(
            b"""<html><body><main><section><h2>Body</h2><p>Prose.</p></section>"""
            b"""<section><h2>References</h2>"""
            b"""<div><div>Smith J (2020) A title.</div><div>Lee A (2019) Other.</div></div>"""
            b"""</section></main></body></html>"""
        )
        contents = parser.parse()
        assert contents.native_ref_strings == [
            "Smith J (2020) A title.",
            "Lee A (2019) Other.",
        ]

    def test_void_elements_stay_inside_the_sentence(self):
        texts = _html_texts(
            b"""<html><body><main><section><h2>S</h2>"""
            b"""<div>Some text with <img alt="ABC"> and <output>42</output> units and end.</div>"""
            b"""<p>Next.</p></section></main></body></html>"""
        )
        # An image is a word boundary, as in <p> text; its alt is not read.
        assert texts == ["Some text with and 42 units and end.", "Next."]

    def test_br_separates_and_hr_flushes_bare_section_text(self):
        texts = _html_texts(
            b"""<html><body><main><section><h2>S</h2>A<br>B<hr>C</section></main></body></html>"""
        )
        assert texts == ["A B", "C"]

    def test_buffered_inline_hyperlink_is_captured(self):
        parser = HtmlParser(
            b"""<html><body><main><section><h2>S</h2>"""
            b"""<div>Direct div text with <a href="https://example.org/x">a link</a>.</div>"""
            b"""</section></main></body></html>"""
        )
        parser.parse()
        assert [url for url, _text, _sec, _idx in parser._pending_url_links] == [
            "https://example.org/x"
        ]


class TestMultiTableFigure:
    def test_figcaption_names_only_the_first_table(self):
        parser = HtmlParser(
            b"""<html><body><article><h2>Results</h2>"""
            b"""<figure><figcaption>Table 1. Demographics</figcaption>"""
            b"""<table><tr><th>A</th></tr><tr><td>1</td></tr></table>"""
            b"""<table><tr><th>B</th></tr><tr><td>2</td></tr></table></figure>"""
            b"""</article></body></html>"""
        )
        contents = parser.parse()
        assert [(t.caption, t.label) for t in contents.tables] == [
            ("Table 1. Demographics", "1"),
            (None, None),
        ]


def _add_note_ref(para, note_id: str) -> None:
    """Anchor a footnote reference on a live python-docx paragraph."""
    from lxml import etree

    run = para._element.find(f"{{{_W}}}r")
    if run is None:
        run = etree.SubElement(para._element, f"{{{_W}}}r")
    ref = etree.SubElement(run, f"{{{_W}}}footnoteReference")
    ref.set(f"{{{_W}}}id", note_id)


def _inject_footnotes(docx_bytes: bytes, notes: dict[str, str]) -> bytes:
    """Add a word/footnotes.xml part carrying *notes* to saved DOCX bytes."""
    import zipfile

    from lxml import etree

    parts: dict[str, bytes] = {}
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zin:
        for name in zin.namelist():
            parts[name] = zin.read(name)
    items = "".join(
        f'<w:footnote w:id="{nid}"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:footnote>'
        for nid, text in notes.items()
    )
    parts["word/footnotes.xml"] = (
        f'<w:footnotes xmlns:w="{_W}">' + items + "</w:footnotes>"
    ).encode()
    cts = parts["[Content_Types].xml"].decode()
    if "footnotes" not in cts:
        cts = cts.replace(
            "</Types>",
            '<Override PartName="/word/footnotes.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument'
            '.wordprocessingml.footnotes+xml"/></Types>',
        )
        parts["[Content_Types].xml"] = cts.encode()
    root = etree.fromstring(parts["word/_rels/document.xml.rels"])
    ns = {"rel": "http://schemas.openxmlformats.org/package/2006/relationships"}
    if not root.findall(
        ".//rel:Relationship[@Target='footnotes.xml']",
        ns,
    ):
        rel = etree.SubElement(root, f"{{{ns['rel']}}}Relationship")
        rel.set("Id", "rFootTest")
        rel.set(
            "Type",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
        )
        rel.set("Target", "footnotes.xml")
        parts["word/_rels/document.xml.rels"] = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, blob in parts.items():
            zout.writestr(name, blob)
    return buf.getvalue()


def _segmented(parser: DocxParser, contents):
    parser.apply_segmentation(
        contents,
        [[entry.text] for entry in parser.assembler.entries if entry.needs_segmentation],
    )
    parser.create_content_sections(contents)
    return contents


class TestDocxTitleDerivedStyles:
    def test_author_and_date_based_on_title_stay_body_text(self):
        # Pandoc/Quarto reference docs base Author/Date/Subtitle on Title;
        # following the chain into "Title" made every author line and the
        # date a level-1 section and swallowed the author names.
        doc = Document()
        doc.add_paragraph("A paper", style="Title")
        author_style = doc.styles.add_style("Author", 1)
        author_style.base_style = doc.styles["Title"]
        doc.add_paragraph("Jane Doe", style="Author")
        date_style = doc.styles.add_style("Date", 1)
        date_style.base_style = doc.styles["Title"]
        doc.add_paragraph("2026-02-22", style="Date")
        doc.add_heading("Method", level=1)
        doc.add_paragraph("Body.")
        parser = DocxParser(_save(doc))
        contents = parser.parse()
        assert [s.header for s in contents.sections] == ["Root", "A paper", "Method"]
        assert contents.detected_title == "A paper"
        texts = [entry.text for entry in parser.assembler.entries]
        assert "Jane Doe" in texts
        assert "2026-02-22" in texts

    def test_style_based_on_heading_still_counts(self):
        # Guard: the Title exclusion must not break custom styles derived
        # from Heading N.
        doc = Document()
        based = doc.styles.add_style("APA Heading 1", 1)
        based.base_style = doc.styles["Heading 1"]
        doc.add_paragraph("Method", style="APA Heading 1")
        doc.add_paragraph("Body.")
        contents = DocxParser(_save(doc)).parse()
        assert [s.header for s in contents.sections] == ["Root", "Method"]


class TestDocxOutlineBodyText:
    def test_outline_level_9_is_body_text(self):
        from lxml import etree

        doc = Document()
        doc.add_heading("Introduction", level=1)
        body = doc.add_paragraph("We recruited eighty participants.")
        ppr = body._element.get_or_add_pPr()
        lvl = etree.SubElement(ppr, f"{{{_W}}}outlineLvl")
        lvl.set(f"{{{_W}}}val", "9")
        parser = DocxParser(_save(doc))
        contents = parser.parse()
        assert [s.header for s in contents.sections] == ["Root", "Introduction"]
        assert "We recruited eighty participants." in [
            entry.text for entry in parser.assembler.entries
        ]

    def test_outline_level_9_on_a_style_is_body_text(self):
        import zipfile

        from lxml import etree

        doc = Document()
        doc.add_paragraph("A body sentence.")
        data = _save(doc)
        parts: dict[str, bytes] = {}
        with zipfile.ZipFile(io.BytesIO(data)) as zin:
            for name in zin.namelist():
                parts[name] = zin.read(name)
        root = etree.fromstring(parts["word/styles.xml"])
        ns = {"w": _W}
        (normal,) = [
            s for s in root.findall(".//w:style", ns) if s.get(f"{{{_W}}}styleId") == "Normal"
        ]
        ppr = normal.find("w:pPr", ns)
        if ppr is None:
            ppr = etree.SubElement(normal, f"{{{_W}}}pPr")
        lvl = etree.SubElement(ppr, f"{{{_W}}}outlineLvl")
        lvl.set(f"{{{_W}}}val", "9")
        parts["word/styles.xml"] = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for name, blob in parts.items():
                zout.writestr(name, blob)
        parser = DocxParser(buf.getvalue())
        contents = parser.parse()
        assert [s.header for s in contents.sections] == ["Root"]
        assert [entry.text for entry in parser.assembler.entries] == ["A body sentence."]

    def test_outline_level_1_means_heading_level_2(self):
        from lxml import etree

        doc = Document()
        para = doc.add_paragraph("Discussion")
        ppr = para._element.get_or_add_pPr()
        lvl = etree.SubElement(ppr, f"{{{_W}}}outlineLvl")
        lvl.set(f"{{{_W}}}val", "1")
        contents = DocxParser(_save(doc)).parse()
        (section,) = [s for s in contents.sections if s.header == "Discussion"]
        assert section.level == 2


class TestDocxSectionWordTitle:
    def test_first_heading_with_a_section_word_stays_the_title(self):
        # "Methods for measuring sleep in older adults" contains the
        # substring alias "methods" (score 0.95) — only an exact-alias hit
        # may reject a first Heading 1 as the title.
        doc = Document()
        doc.add_heading("Methods for measuring sleep in older adults", level=1)
        doc.add_paragraph("Body text.")
        contents = DocxParser(_save(doc)).parse()
        assert contents.detected_title == "Methods for measuring sleep in older adults"

    def test_first_heading_that_is_a_section_is_not_the_title(self):
        doc = Document()
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph("Body text.")
        contents = DocxParser(_save(doc)).parse()
        assert contents.detected_title is None

    def test_first_heading_after_hand_formatted_lines_keeps_mains_reading(self):
        # An APA manuscript: hand-formatted title lines, then a Heading 1
        # "Author Note" (a substring alias only) holding the article's DOI.
        # Main typed that block TITLE, which keeps it in the front matter the
        # DOI selection reads; refusing the title after any earlier content
        # lost the DOI in the no-LLM export, so that guard is not applied.
        doc = Document()
        doc.add_paragraph("Attention and Memory in Older Adults")
        doc.add_paragraph("Jane Doe and Alice Roe")
        doc.add_heading("Author Note", level=1)
        doc.add_paragraph(
            "The final article is available, upon publication, at: "
            "https://doi.org/10.1037/xge0001234"
        )
        doc.add_heading("Abstract", level=1)
        doc.add_paragraph("We studied attention.")
        contents = DocxParser(_save(doc)).parse()
        assert contents.detected_title == "Author Note"


class TestDocxUnpairedCaptions:
    _LINKED_DRAWING = (
        '<w:drawing xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        "<a:graphic><a:graphicData><pic:pic><pic:blipFill>"
        '<a:blip r:link="rId9"/>'
        "</pic:blipFill></pic:pic></a:graphicData></a:graphic></w:drawing>"
    )

    def test_linked_picture_caption_stays_body_text(self):
        doc = Document()
        doc.add_paragraph("Intro.")
        _raw_paragraph(doc, f"<w:r><w:t>See the image: </w:t></w:r>{self._LINKED_DRAWING}")
        doc.add_paragraph("Figure 1. A linked picture.", style="Caption")
        parser = DocxParser(_save(doc))
        contents = parser.parse()
        assert contents.figures == []
        assert "Figure 1. A linked picture." in [entry.text for entry in parser.assembler.entries]

    def test_picture_inside_a_heading_keeps_its_caption(self):
        from docx.shared import Pt

        doc = Document()
        pic_para = doc.add_paragraph()
        pic_para.add_run().add_picture(io.BytesIO(_PNG), width=Pt(10))
        pic_para.style = doc.styles["Heading 1"]
        doc.add_paragraph("Figure 1. Picture in a heading paragraph.", style="Caption")
        parser = DocxParser(_save(doc))
        contents = parser.parse()
        assert contents.figures == []
        assert "Figure 1. Picture in a heading paragraph." in [
            entry.text for entry in parser.assembler.entries
        ]


class TestDocxMergedCellFootnotes:
    def test_footnote_in_horizontally_merged_cell_queues_once(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=3)
        table.cell(0, 0).merge(table.cell(0, 2))
        cell = table.cell(0, 0)
        cell.paragraphs[0].add_run("Merged cell.")
        _add_note_ref(cell.paragraphs[0], "1")
        data = _inject_footnotes(_save(doc), {"1": "The only note."})
        parser = DocxParser(data)
        contents = _segmented(parser, parser.parse())
        headers = [s.header for s in contents.sections]
        assert headers.count("Footnote 1") == 1
        assert "Footnote 2" not in headers

    def test_footnote_in_vertically_merged_cell_queues_once(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).merge(table.cell(1, 0))
        cell = table.cell(0, 0)
        cell.paragraphs[0].add_run("Merged cell.")
        _add_note_ref(cell.paragraphs[0], "1")
        data = _inject_footnotes(_save(doc), {"1": "The only note."})
        parser = DocxParser(data)
        contents = _segmented(parser, parser.parse())
        headers = [s.header for s in contents.sections]
        assert headers.count("Footnote 1") == 1
        assert "Footnote 2" not in headers

    def test_footnote_in_table_cell_is_emitted(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Body cell."
        _add_note_ref(table.cell(0, 0).paragraphs[0], "1")
        data = _inject_footnotes(_save(doc), {"1": "A cell note."})
        parser = DocxParser(data)
        contents = _segmented(parser, parser.parse())
        assert "Footnote 1" in [s.header for s in contents.sections]
        assert "A cell note." in [s.text for s in contents.sentences]


class TestDocxCellCitationsAndCaptionNotes:
    def test_sdt_citation_in_a_table_cell_is_kept(self):
        from lxml import etree

        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Effect"
        table.cell(0, 1).text = "0.5"
        table.cell(1, 0).text = "Group"
        target = table.cell(1, 1)
        target.text = ""
        frag = etree.fromstring(
            f'<w:p xmlns:w="{_W}">'
            "<w:r><w:t>Effect </w:t></w:r>"
            '<w:sdt><w:sdtPr><w:alias w:val="citation"/></w:sdtPr>'
            "<w:sdtContent><w:r><w:t>(Smith, 2020)</w:t></w:r></w:sdtContent></w:sdt>"
            "</w:p>"
        )
        para_el = target.paragraphs[0]._element
        for child in list(frag):
            para_el.append(child)
        contents = DocxParser(_save(doc)).parse()
        (kept,) = contents.tables
        assert kept.df.iloc[0, 1] == "Effect (Smith, 2020)"

    def test_footnote_on_a_table_caption_is_emitted_in_order(self):
        doc = Document()
        caption = doc.add_paragraph("Table 1. Demographics", style="Caption")
        _add_note_ref(caption, "1")
        table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "Group"
        table.cell(1, 0).text = "A"
        body = doc.add_paragraph("Body note here.")
        _add_note_ref(body, "2")
        data = _inject_footnotes(_save(doc), {"1": "Sourced from the lab.", "2": "A body note."})
        parser = DocxParser(data)
        contents = _segmented(parser, parser.parse())
        notes = [
            s.text
            for s in contents.sentences
            if s.text in ("Sourced from the lab.", "A body note.")
        ]
        assert notes == ["Sourced from the lab.", "A body note."]
        assert [x.contents for x in contents.xrefs if x.xref_type == "foot"] == ["1", "2"]


class TestDocxSymbolChars:
    def test_non_letter_symbol_chars_read_as_text(self):
        doc = Document()
        _raw_paragraph(
            doc,
            "<w:r><w:t>p </w:t></w:r>"
            '<w:r><w:sym w:font="Symbol" w:char="F0B1"/></w:r>'
            "<w:r><w:t> .05, M </w:t></w:r>"
            '<w:r><w:sym w:font="Symbol" w:char="F061"/></w:r>'
            "<w:r><w:t> SD, 2 </w:t></w:r>"
            '<w:r><w:sym w:font="Symbol" w:char="F0A3"/></w:r>'
            "<w:r><w:t> 3, 20</w:t></w:r>"
            '<w:r><w:sym w:font="Symbol" w:char="F0B0"/></w:r>'
            "<w:r><w:t> and x </w:t></w:r>"
            '<w:r><w:sym w:font="Symbol" w:char="F0B4"/></w:r>'
            "<w:r><w:t> y, see </w:t></w:r>"
            '<w:r><w:sym w:font="Symbol" w:char="F0AE"/></w:r>'
            "<w:r><w:t> note, a </w:t></w:r>"
            '<w:r><w:sym w:font="Symbol" w:char="F0B3"/></w:r>'
            "<w:r><w:t> b, c </w:t></w:r>"
            '<w:r><w:sym w:font="Symbol" w:char="F0B9"/></w:r>'
            "<w:r><w:t> d.</w:t></w:r>",
        )
        _, texts = _docx_texts(_save(doc))
        assert texts == ["p ± .05, M \u03b1 SD, 2 ≤ 3, 20° and x × y, see → note, a ≥ b, c ≠ d."]


class TestDocxPlaceholders:
    _PLACEHOLDER_SDT = (
        '<w:sdt><w:sdtPr><w:alias w:val="Author"/><w:showingPlcHdr/></w:sdtPr>'
        "<w:sdtContent><w:r><w:t>Click or tap here to enter text.</w:t></w:r></w:sdtContent></w:sdt>"
    )

    def test_inline_placeholder_prompt_is_skipped(self):
        doc = Document()
        _raw_paragraph(
            doc,
            "<w:r><w:t>Author: </w:t></w:r>" + self._PLACEHOLDER_SDT,
        )
        _, texts = _docx_texts(_save(doc))
        assert texts == ["Author:"]

    def test_block_placeholder_prompt_is_skipped(self):
        import zipfile

        from lxml import etree

        doc = Document()
        doc.add_paragraph("Real paragraph.")
        data = _save(doc)
        parts: dict[str, bytes] = {}
        with zipfile.ZipFile(io.BytesIO(data)) as zin:
            for name in zin.namelist():
                parts[name] = zin.read(name)
        root = etree.fromstring(parts["word/document.xml"])
        ns = {"w": _W}
        (body,) = root.findall(".//w:body", ns)
        sdt = etree.fromstring(
            f'<w:sdt xmlns:w="{_W}">'
            '<w:sdtPr><w:alias w:val="Title"/><w:showingPlcHdr/></w:sdtPr>'
            "<w:sdtContent><w:p><w:r><w:t>Click or tap here to enter text.</w:t></w:r></w:p></w:sdtContent>"
            "</w:sdt>"
        )
        body.append(sdt)
        parts["word/document.xml"] = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for name, blob in parts.items():
                zout.writestr(name, blob)
        _, texts = _docx_texts(buf.getvalue())
        assert texts == ["Real paragraph."]


class TestDocxVmlTextbox:
    def test_vml_text_box_paragraphs_keep_word_boundaries(self):
        doc = Document()
        _raw_paragraph(
            doc,
            '<w:r><w:t>Main text.</w:t></w:r><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml">'
            "<v:textbox><w:txbxContent>"
            "<w:p><w:r><w:t>Box heading</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>Box body.</w:t></w:r></w:p>"
            "</w:txbxContent></v:textbox></v:shape></w:pict>",
        )
        _, texts = _docx_texts(_save(doc))
        assert texts == ["Main text. Box heading Box body."]

    def test_mid_paragraph_drawing_box_keeps_boundaries(self):
        # Guard (mutant M23): a text box between runs must separate words,
        # not fuse them.
        doc = Document()
        _raw_paragraph(
            doc,
            '<w:r><w:t>Body sentence. </w:t><mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
            "<mc:Choice><w:drawing><w:txbxContent>"
            "<w:p><w:r><w:t>Boxed</w:t></w:r><w:r><w:t> pull quote.</w:t></w:r></w:p>"
            "</w:txbxContent></w:drawing></mc:Choice></mc:AlternateContent></w:r>"
            "<w:r><w:t>After.</w:t></w:r>",
        )
        _, texts = _docx_texts(_save(doc))
        assert texts == ["Body sentence. Boxed pull quote. After."]


class TestDocxRuby:
    def test_ruby_reading_is_not_duplicated(self):
        doc = Document()
        _raw_paragraph(
            doc,
            "<w:r><w:t>See </w:t></w:r>"
            "<w:ruby><w:rt><w:r><w:t>reading</w:t></w:r></w:rt>"
            "<w:rubyBase><w:r><w:t>base</w:t></w:r></w:rubyBase></w:ruby>"
            "<w:r><w:t> here.</w:t></w:r>",
        )
        _, texts = _docx_texts(_save(doc))
        assert texts == ["See base here."]


class TestDocxTableCaptionPairing:
    def _table_doc(self, build):
        doc = Document()
        build(doc)
        return _save(doc)

    def _add_pic(self, doc):
        from docx.shared import Pt

        doc.add_paragraph().add_run().add_picture(io.BytesIO(_PNG), width=Pt(10))

    def test_table_between_picture_and_caption_is_skipped(self):
        # A float may sit between a picture and its caption.
        def build(doc):
            self._add_pic(doc)
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "x"
            doc.add_paragraph("Figure 1. Under a table.", style="Caption")

        contents = DocxParser(self._table_doc(build)).parse()
        assert [fig.caption for fig in contents.figures] == ["Figure 1. Under a table."]

    def test_table_claimed_caption_is_not_repaired(self):
        # The table pairing runs first: its caption never re-pairs with a
        # neighbouring picture.
        def build(doc):
            self._add_pic(doc)
            doc.add_paragraph("Table 1. Tabular data.", style="Caption")
            table = doc.add_table(rows=2, cols=1)
            table.cell(0, 0).text = "A"
            table.cell(1, 0).text = "1"

        contents = DocxParser(self._table_doc(build)).parse()
        assert [fig.caption for fig in contents.figures] == [None]
        assert [tbl.caption for tbl in contents.tables] == ["Table 1. Tabular data."]


class TestNestedHiddenChrome:
    """Noise removal must survive hidden containers that hold child elements.

    Decomposing a matched tag while iterating find_all left its descendants
    dead (attrs None); the next tag.get() raised AttributeError, so every
    eLife page (aria-hidden separators with children, visuallyhidden spans
    around counters) failed to parse.
    """

    def test_aria_hidden_element_with_a_child(self):
        texts = _html_texts(
            b"""<html><body><article><h2>Intro</h2>"""
            b"""<p>Text <span aria-hidden="true"><i>x</i></span> here.</p></article></body></html>"""
        )
        assert texts == ["Text here."]

    def test_visuallyhidden_element_with_a_child(self):
        texts = _html_texts(
            b"""<html><body><article><h2>Intro</h2>"""
            b"""<p>Text <span class="visuallyhidden">(<span class="n">0</span> notes)</span>"""
            b""" here.</p></article></body></html>"""
        )
        assert texts == ["Text here."]

    def test_nested_landmark_roles(self):
        texts = _html_texts(
            b"""<html><body><div role="navigation"><div role="search">Find</div></div>"""
            b"""<article><h2>Intro</h2><p>Body.</p></article></body></html>"""
        )
        assert texts == ["Body."]


class TestPublisherFloatFurniture:
    """eLife-shaped figure header strips and asset links are not prose."""

    _PAGE = (
        b"""<html><body><main><section><h2>Results</h2>"""
        b"""<p class="paragraph">We saw an effect (Figure 2).</p>"""
        b"""<div id="fig2" class="asset-viewer-inline">"""
        b"""<div class="asset-viewer-inline__header_panel">"""
        b"""<div class="asset-viewer-inline__header_text">"""
        b"""<span class="asset-viewer-inline__header_text__prominent">Figure 2</span>"""
        b""" with 2 supplements <a href="/articles/1/figures#fig2">see all</a></div>"""
        b"""<div class="asset-viewer-inline__figure_access">"""
        b"""<a href="https://example.org/fig2.jpg" download="Download">Download</a></div></div>"""
        b"""<figure class="captioned-asset"><img src="fig2.jpg">"""
        b"""<figcaption>Effect sizes by group.</figcaption></figure></div>"""
        b"""<div class="asset-viewer-inline__header_text"><span>Table 1</span></div>"""
        b"""<div class="asset-viewer-inline__header_text">Author response image 1</div>"""
        b"""<div class="asset-viewer-inline__header_text">Appendix 1\xe2\x80\x94figure 2</div>"""
        b"""<div><a href="https://example.org/f1-data1.xlsx">Download elife-1-fig1-data1-v1.xlsx</a></div>"""
        b"""<div class="math-block"><math><mi>x</mi><mo>=</mo><mn>2</mn></math></div>"""
        b"""<div class="message-bar">The following data sets were generated</div>"""
        b"""<p>Figure 3</p>"""
        b"""</section></main></body></html>"""
    )

    def test_labels_supplement_toggles_and_download_links_are_dropped(self):
        texts = _html_texts(self._PAGE)
        assert texts == [
            "We saw an effect (Figure 2).",
            "x=2",
            "The following data sets were generated",
            # Guard: a <p> is read exactly as before, even when it is a label.
            "Figure 3",
        ]

    def test_figure_is_still_recorded(self):
        contents = HtmlParser(self._PAGE).parse()
        assert [fig.caption for fig in contents.figures] == ["Effect sizes by group."]

    def test_prose_that_starts_with_a_float_word_is_kept(self):
        texts = _html_texts(
            b"""<html><body><main><section><h2>S</h2>"""
            b"""<div>Figure 1 shows the effect.</div><div>Table salt was used.</div>"""
            b"""</section></main></body></html>"""
        )
        assert texts == ["Figure 1 shows the effect.", "Table salt was used."]


class TestCitationBlocks:
    def test_citation_block_with_a_paragraph_is_still_read(self):
        # Guard: only a block with nothing base read as body text is skipped.
        texts = _html_texts(
            b"""<html><body><main><section><h2>S</h2>"""
            b"""<div class="citation"><p>Cite as: Smith (2020).</p><span>(2020)</span></div>"""
            b"""</section></main></body></html>"""
        )
        assert texts == ["Cite as: Smith (2020)."]

    def test_generic_list_inside_a_citation_block_is_still_read(self):
        texts = _html_texts(
            b"""<html><body><main><section><h2>S</h2><p>Prose.</p>"""
            b"""<div class="citation"><ul><li>First point.</li></ul></div>"""
            b"""</section></main></body></html>"""
        )
        assert texts == ["Prose.", "First point."]


class TestInlineImageAltText:
    def test_icon_alt_text_is_not_read_into_an_author_item(self):
        texts = _html_texts(
            b"""<html><body><main><section><h2>Authors</h2><ol>"""
            b"""<li><a href="#a1">Jon Clardy</a>&nbsp;<picture>"""
            b"""<img src="icon.png" alt="Is a corresponding author"></picture></li>"""
            b"""</ol></section></main></body></html>"""
        )
        assert texts == ["Jon Clardy"]


class TestDocxLaterRowCellFootnotes:
    """Cell dedupe must compare elements, not reusable proxy ids.

    python-docx builds new cell proxies for every row.cells call; the freed
    proxies' ids come back for the next row, so an id() set skipped cells in
    every row after the first and their notes were lost.
    """

    def _notes(self, doc, notes):
        data = _inject_footnotes(_save(doc), notes)
        parser = DocxParser(data)
        parser.parse()
        return [text for text, _sec, _idx, _kind in parser._pending_footnotes]

    def test_footnote_in_a_second_row_cell(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=3)
        for col, head in enumerate(("h1", "h2", "h3")):
            table.cell(0, col).text = head
        table.cell(1, 0).text = "a"
        table.cell(1, 1).text = "b"
        _add_note_ref(table.cell(1, 1).paragraphs[0], "1")
        table.cell(1, 2).text = "c"
        assert self._notes(doc, {"1": "Note one."}) == ["Note one."]

    def test_footnotes_in_every_cell_of_the_third_row(self):
        doc = Document()
        table = doc.add_table(rows=3, cols=3)
        for row in range(3):
            for col in range(3):
                table.cell(row, col).text = f"r{row}c{col}"
        for col, note in enumerate(("1", "2", "3")):
            _add_note_ref(table.cell(2, col).paragraphs[0], note)
        notes = {"1": "Note one.", "2": "Note two.", "3": "Note three."}
        assert self._notes(doc, notes) == ["Note one.", "Note two.", "Note three."]

    def test_footnote_in_a_vertical_merge_below_a_header_row_queues_once(self):
        doc = Document()
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "Group"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).merge(table.cell(2, 0))
        table.cell(1, 0).paragraphs[0].add_run("Merged.")
        _add_note_ref(table.cell(1, 0).paragraphs[0], "1")
        table.cell(1, 1).text = "1"
        table.cell(2, 1).text = "2"
        assert self._notes(doc, {"1": "The only note."}) == ["The only note."]


class TestDocxSymbolEncoding:
    """w:sym Symbol-font codes follow the font's built-in encoding."""

    def _text(self, char: str) -> str:
        doc = Document()
        _raw_paragraph(
            doc,
            "<w:r><w:t>a</w:t></w:r>"
            f'<w:r><w:sym w:font="Symbol" w:char="{char}"/></w:r>'
            "<w:r><w:t>b</w:t></w:r>",
        )
        _parser, texts = _docx_texts(_save(doc))
        return texts[0]

    def test_angle_bracket_and_integral(self):
        assert self._text("F0F1") == "a⟩b"
        assert self._text("F0F2") == "a∫b"

    def test_bracket_pieces_are_not_read_as_characters(self):
        # F0F9/F0FA are right-bracket pieces, not an angle bracket and an
        # integral.
        assert self._text("F0F9") == "ab"
        assert self._text("F0FA") == "ab"

    def test_shared_ascii_and_logic_symbols(self):
        assert self._text("F03D") == "a=b"
        assert self._text("F03C") == "a<b"
        assert self._text("F05C") == "a∴b"


class TestReferenceSectionLeadIns:
    def test_digitless_lead_in_is_not_a_reference(self):
        # eLife's second dataset list follows the first, which already
        # opened a References section; its lead-in is not a citation.
        parser = HtmlParser(
            b"""<html><body><main><section><h2>Data availability</h2><p>Deposited.</p>"""
            b"""<div class="message-bar">The following data sets were generated</div>"""
            b"""<ol class="reference-list"><li>Smith J (2020) Data one. GEO GSE1.</li></ol>"""
            b"""<div class="message-bar">The following previously published data sets were used</div>"""
            b"""<ol class="reference-list"><li>Lee A (2019) Data two. GEO GSE2.</li></ol>"""
            b"""<div>Doe K (2018) A flat div reference. J 1:2.</div>"""
            b"""</section></main></body></html>"""
        )
        contents = parser.parse()
        assert contents.native_ref_strings == [
            "Smith J (2020) Data one. GEO GSE1.",
            "Lee A (2019) Data two. GEO GSE2.",
            "Doe K (2018) A flat div reference. J 1:2.",
        ]
        assert "The following data sets were generated" in [
            entry.text for entry in parser.assembler.entries
        ]


class TestNavigationLists:
    """Lists of bare page-navigation links are chrome, not body text.

    Base opened a References section at the "cite this article" block, so
    the copy/download buttons and the page's tab navigation after it were
    reference strings; with the block skipped they must not become body
    sentences either.
    """

    def test_button_download_and_jump_link_lists_are_dropped(self):
        parser = HtmlParser(
            b"""<html><body><main><section><h6>Cite this article</h6><p>Prose.</p>"""
            b"""<div class="reference"><ol class="reference__authors_list">"""
            b"""<li>Anmo J Kim</li></ol><div class="reference__title">T</div></div>"""
            b"""<ol class="button-collection"><li><button>Copy to clipboard</button></li>"""
            b"""<li><a href="/a.bib" class="button button--secondary">Download BibTeX</a></li></ol>"""
            b"""<ul class="article-download-list"><li><a href="/a.bib">Download BibTeX</a></li>"""
            b"""<li><a href="/a.ris">Download .RIS</a></li></ul>"""
            b"""<ul class="view-selector__list"><li><a href="/a#content">Article</a></li>"""
            b"""<li></li><li><a href="#abstract">Abstract</a></li>"""
            b"""<li><a href="#s1">Introduction</a></li></ul>"""
            b"""<p>Insight body text without a heading.</p></section>"""
            b"""<section><h2>Abstract</h2><p>Abs.</p></section>"""
            b"""<section><h2>Introduction</h2><p>Intro.</p></section>"""
            b"""</main></body></html>"""
        )
        contents = parser.parse()
        assert [e.text for e in parser.assembler.entries] == [
            "Prose.",
            "Insight body text without a heading.",
            "Abs.",
            "Intro.",
        ]
        assert (contents.native_ref_strings or []) == []

    def test_author_and_resource_link_lists_stay(self):
        texts = _html_texts(
            b"""<html><body><main><section><h2>S</h2><p>Prose.</p>"""
            b"""<ol class="author_list"><li><a href="/articles/1#x1">Jon Clardy</a></li>"""
            b"""<li><a href="/articles/1#x2">Nicole King</a></li></ol>"""
            b"""<ul><li><a href="https://osf.io/abc">https://osf.io/abc</a></li></ul>"""
            b"""<ul><li>See <a href="#fig1">Figure 1</a> for details.</li></ul>"""
            b"""</section></main></body></html>"""
        )
        assert texts == [
            "Prose.",
            "Jon Clardy",
            "Nicole King",
            "https://osf.io/abc",
            "See Figure 1 for details.",
        ]


class TestStructuredPageDescription:
    def test_description_is_the_abstract_of_a_structured_editorial(self):
        # An eLife editorial carries dc.title and a DOI and no abstract meta;
        # its description is its JATS abstract, which main exported.
        meta = (
            HtmlParser(
                b"""<html><head><meta name="dc.title" content="Editorial">"""
                b"""<meta name="dc.identifier" content="doi:10.7554/eLife.00855">"""
                b"""<meta name="description" content="It is time to rethink assessment.">"""
                b"""</head><body><main><h2>Body</h2><p>Text.</p></main></body></html>"""
            )
            .parse()
            .preparsed_metadata
        )
        assert meta.abstract == "It is time to rethink assessment."

    def test_abstract_meta_still_wins_over_the_description(self):
        meta = (
            HtmlParser(
                b"""<html><head><meta name="dc.title" content="Paper">"""
                b"""<meta name="dc.identifier" content="doi:10.7554/eLife.1">"""
                b"""<meta name="dc.description" content="The digest.">"""
                b"""<meta name="description" content="Impact statement.">"""
                b"""</head><body><main><h2>Body</h2><p>Text.</p></main></body></html>"""
            )
            .parse()
            .preparsed_metadata
        )
        assert meta.abstract == "The digest."

    def test_description_is_not_the_abstract_when_the_page_prints_one(self):
        # A research article's description is its impact statement; the
        # printed Abstract section must stay the abstract.
        meta = (
            HtmlParser(
                b"""<html><head><meta name="dc.title" content="Paper">"""
                b"""<meta name="dc.identifier" content="doi:10.7554/eLife.03600">"""
                b"""<meta name="description" content="Building on previous work, we used it.">"""
                b"""</head><body><main><h2>Abstract</h2><p>MicroED is a method.</p>"""
                b"""</main></body></html>"""
            )
            .parse()
            .preparsed_metadata
        )
        assert meta.abstract == ""
