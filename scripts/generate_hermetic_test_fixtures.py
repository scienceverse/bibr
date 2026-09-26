"""Generate the committed synthetic fixtures for the hermetic regression tests.

Writes three tiny deterministic files under ``tests/fixtures/`` with no third
party PDF/Office writers (only the stdlib, plus ``python-docx`` for the .docx
shell):

- ``native_text_bbox_bleed_sample.pdf``: two Helvetica lines where a region
  bbox top edge slices through the upper line's glyph boxes while their
  centers stay outside. Under the old intersect rule the query bleeds the
  upper line's descender fragments (``pyy``); center-containment returns the
  lower line exactly. Used by ``tests/test_native_text.py``.
- ``ref_geometry_hanging_indent_sample.pdf``: a ``References`` header plus 22
  hanging-indent reference lines (starts at x=72, continuations at x=90).
  Used by ``tests/ocr/test_ref_geometry.py`` instead of the uncommitted gold
  PDF.
- ``footnotes_sample.docx``: two real footnotes plus separator
  pseudo-notes. Used by ``tests/input/test_docx_footnotes.py``.

Usage: ``python scripts/generate_hermetic_test_fixtures.py`` from the repo root.
"""

from __future__ import annotations

import argparse
import io
import sys
import zipfile
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
FIXTURES = REPO_ROOT / "tests" / "fixtures"

PAGE_W, PAGE_H, FONT_SIZE = 612.0, 792.0, 12.0
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _esc(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_pdf(lines: list[tuple[float, float, str]]) -> bytes:
    """Minimal single-page PDF; ``lines`` are ``(x, baseline_y, text)``."""
    content = "".join(
        f"BT /F1 {FONT_SIZE:g} Tf {x:g} {y:g} Td ({_esc(t)}) Tj ET\n" for x, y, t in lines
    )
    encoded = content.encode("latin-1")
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {PAGE_W:g} {PAGE_H:g}] "
        "/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(encoded)} >>\nstream\n{content}endstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode("latin-1")
    out += b"".join(f"{o:010d} 00000 n \n".encode("latin-1") for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode(
        "latin-1"
    )
    return bytes(out)


def write_bbox_bleed_pdf(path: Path) -> None:
    path.write_bytes(
        build_pdf(
            [
                (72, 700, "Department of Synthetic Medicine University"),
                (72, 682, "Second line of body text content here"),
            ]
        )
    )


def write_ref_geometry_pdf(path: Path) -> None:
    lines: list[tuple[float, float, str]] = [("References", 72)]
    for i in range(11):
        lines.append((f"Smith, A. ({2000 + i}). Synthetic study number {i} in medicine.", 72))
        lines.append((f"continued discussion of that synthetic study number {i} here.", 90))
    ordered = [(x, 740 - 14 * k, text) for k, (text, x) in enumerate(lines)]
    path.write_bytes(build_pdf(ordered))


def write_footnotes_docx(path: Path) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    doc.add_paragraph("Synthetic body text for the footnotes fixture.")
    para = doc.add_paragraph("A sentence with a footnote")
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), "2")
    para.add_run()._r.append(ref)
    para.add_run(" and a second footnote")
    ref2 = OxmlElement("w:footnoteReference")
    ref2.set(qn("w:id"), "3")
    para.runs[-1]._r.append(ref2)
    para.add_run(".")

    buf = io.BytesIO()
    doc.save(buf)
    payload = buf.getvalue()

    footnotes = (
        f'<w:footnotes xmlns:w="{W_NS}">'
        '<w:footnote w:id="0" w:type="separator"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:id="1" w:type="continuationSeparator">'
        "<w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>"
        '<w:footnote w:id="2"><w:p><w:r><w:t>First synthetic footnote text.</w:t></w:r></w:p></w:footnote>'
        '<w:footnote w:id="3"><w:p><w:r><w:t>Second synthetic footnote with </w:t></w:r>'
        "<w:r><w:t>split runs.</w:t></w:r></w:p></w:footnote>"
        "</w:footnotes>"
    ).encode()
    rels_extra = (
        b'<Relationship Id="rId99" '
        b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" '
        b'Target="footnotes.xml"/>'
    )
    override = (
        b'<Override PartName="/word/footnotes.xml" '
        b'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
    )

    src = zipfile.ZipFile(io.BytesIO(payload))
    dst_buf = io.BytesIO()
    # Fixed entry timestamps: python-docx stamps "now" into every entry, so
    # without this the committed bytes would differ on every regeneration.
    stamp = (2020, 1, 2, 0, 0, 0)
    with zipfile.ZipFile(dst_buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"</Relationships>", rels_extra + b"</Relationships>")
            elif item.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", override + b"</Types>")
            info = zipfile.ZipInfo(item.filename, date_time=stamp)
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, data)
        info = zipfile.ZipInfo("word/footnotes.xml", date_time=stamp)
        info.compress_type = zipfile.ZIP_DEFLATED
        dst.writestr(info, footnotes)
    path.write_bytes(dst_buf.getvalue())


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=FIXTURES)
    args = parser.parse_args(argv)
    args.out.mkdir(parents=True, exist_ok=True)
    write_bbox_bleed_pdf(args.out / "native_text_bbox_bleed_sample.pdf")
    write_ref_geometry_pdf(args.out / "ref_geometry_hanging_indent_sample.pdf")
    write_footnotes_docx(args.out / "footnotes_sample.docx")
    for name in (
        "native_text_bbox_bleed_sample.pdf",
        "ref_geometry_hanging_indent_sample.pdf",
        "footnotes_sample.docx",
    ):
        size = (args.out / name).stat().st_size
        print(f"{name}: {size} bytes")
    return 0


if __name__ == "__main__":
    sys.exit(main())
